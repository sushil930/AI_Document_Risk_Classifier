from pathlib import Path
import io
import zipfile
import pymupdf as fitz
from docx import Document as WordDocument

def extract_text(upload):
    """Extract TXT, DOCX, or PDF text and retain page or paragraph references."""
    suffix = Path(upload.name).suffix.lower().lstrip(".")
    upload.seek(0)
    if suffix not in {"pdf", "docx", "txt"}: raise ValueError("Upload a PDF, DOCX, or TXT file.")
    raw = upload.read()
    if suffix == "pdf":
        if not raw.lstrip().startswith(b"%PDF-"): raise ValueError("File contents do not match the PDF extension.")
        pdf = fitz.open(stream=raw, filetype="pdf")
        text = "\n\n".join(f"[Page {i + 1}]\n{page.get_text()}" for i, page in enumerate(pdf))
    elif suffix == "docx":
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if "word/document.xml" not in archive.namelist(): raise ValueError("Invalid DOCX content.")
        except zipfile.BadZipFile as exc: raise ValueError("Invalid DOCX content.") from exc
        doc = WordDocument(io.BytesIO(raw))
        text = "\n".join(f"[Paragraph {i + 1}] {p.text}" for i, p in enumerate(doc.paragraphs) if p.text.strip())
        for ti, table in enumerate(doc.tables, 1):
            text += "\n" + "\n".join(f"[Table {ti}, Row {ri + 1}] " + " | ".join(c.text for c in row.cells) for ri, row in enumerate(table.rows))
    else:
        if b"\x00" in raw: raise ValueError("TXT file contains binary data.")
        try: text = raw.decode("utf-8-sig")
        except UnicodeDecodeError as exc: raise ValueError("TXT content is not valid UTF-8 text.") from exc
    return text.strip(), suffix
